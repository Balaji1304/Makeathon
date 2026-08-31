"""Generate German DOCX from PROJECT_DOCUMENTATION.md using python-docx. Black & white, names replaced."""
import re
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

OUTPUT_PATH = r"D:\projects\Makeathon\docs\GreenTrack_Projektdokumentation.docx"

# --- German Translation Map ---
TRANSLATIONS = {
    "GreenTrack Logistics Control Tower": "GreenTrack Logistik-Kontrollturm",
    "Enterprise Data-Driven Supply Chain & Carbon Optimization Platform": "Unternehmensplattform fuer datengetriebene Supply-Chain- und CO2-Optimierung",
    "Makeathon Team": "Balaji Loganathan",
    "August 2026": "August 2026",
    "Table of Contents": "Inhaltsverzeichnis",
    "1. DESCRIPTION": "1. BESCHREIBUNG",
    "1.1 Participants": "1.1 Teilnehmer",
    "1.2 What Problem Should Be Solved": "1.2 Zu loesendes Problem",
    "1.3 Software Functionalities": "1.3 Softwarefunktionen",
    "2. USER STORIES": "2. BENUTZERGESCHICHTEN",
    "2.1 Profiles and Roles": "2.1 Profile und Rollen",
    "2.2 Dashboard & Analytics": "2.2 Dashboard & Analytik",
    "2.3 Fleet Management": "2.3 Flottenverwaltung",
    "2.4 Order & Route Planning": "2.4 Auftrags- und Routenplanung",
    "2.5 Alert & Recommendation Engine": "2.5 Warn- und Empfehlungssystem",
    "2.6 ML Prescriptive Simulator": "2.6 ML-Simulator fuer prskriptive Analytik",
    "2.7 Interactive Map Visualization": "2.7 Interaktive Kartenvisualisierung",
    "2.8 Profile & Settings": "2.8 Profil & Einstellungen",
    "3. PLANNING AND DEVELOPMENT": "3. PLANUNG UND ENTWICKLUNG",
    "3.1 Estimation": "3.1 Aufwandsschaetzung",
    "3.2 Milestones": "3.2 Meilensteine",
    "4. FINAL PRODUCT": "4. ENDPRODUKT",
    "4.1 Installation and Configuration": "4.1 Installation und Konfiguration",
    "4.2 First Test": "4.2 Erster Test",
    "4.3 Database Management": "4.3 Datenbankverwaltung",
    "4.4 User Guide": "4.4 Benutzerhandbuch",
    "5. STRUCTURE": "5. STRUKTUR",
    "5.1 Architecture": "5.1 Architektur",
    "5.2 Data Pipeline": "5.2 Datenpipelinie",
    "5.3 ML Pipeline": "5.3 ML-Pipelinie",
    "6. REVIEW": "6. REVIEWS",
    "6.1 Work Distribution": "6.1 Arbeitsverteilung",
    "6.2 Time Distribution": "6.2 Zeiterfassung",
    "6.3 Achieved and Unachieved User Stories": "6.3 Erreichte und nicht erreichte Benutzergeschichten",
    "6.4 Problems": "6.4 Probleme",
    "6.5 Future Improvements": "6.5 Zukunftige Verbesserungen",
    "6.6 Software Used": "6.6 Verwendete Software",
}

SECTION_DE = {
    "1. BESCHREIBUNG": "BESCHREIBUNG",
    "2. BENUTZERGESCHICHTEN": "BENUTZERGESCHICHTEN",
    "3. PLANUNG UND ENTWICKLUNG": "PLANUNG UND ENTWICKLUNG",
    "4. ENDPRODUKT": "ENDPRODUKT",
    "5. STRUKTUR": "STRUKTUR",
    "6. REVIEWS": "REVIEWS",
}

def translate_line(line):
    for en, de in TRANSLATIONS.items():
        if en in line:
            line = line.replace(en, de)
    return line

def translate_full(text):
    text = text.replace("Premanathan Aarthi Manivannan", "Balaji Loganathan")
    text = text.replace("Jeyanth Shanmugasundaram", "Balaji Loganathan")
    text = text.replace("Rahul Suresh", "Balaji Loganathan")
    text = text.replace("Premanathan", "Balaji Loganathan")
    text = text.replace("Jeyanth", "Balaji Loganathan")
    text = text.replace("Rahul", "Balaji Loganathan")
    for en, de in TRANSLATIONS.items():
        text = text.replace(en, de)
    return text

# --- German content replacements ---
CONTENT_DE = {
    # Section 1.2
    "In modern supply chains, organizations struggle with two critical inefficiencies: **Fragmented Data Visibility** and **High Carbon Footprints**. Dispatchers often rely on outdated systems that send heavy, underutilized combustion-engine vehicles on routes better suited for light electric vans.":
    "In modernen Supply Chains kaempfen Unternehmen mit zwei kritischen Ineffizienzen: **Fragmentierte Datensichtbarkeit** und **hohe CO2-Belastung**. Disponenten verlassen sich haeufig auf veraltete Systeme, die schwere, unterausgelastete Verbrennungsfahrzeuge auf Routen schicken, die besser fuer leichte Elektrovan geeignet waeren.",
    "A software platform will be developed to provide a centralized, data-driven **Logistics Control Tower** that ingests raw operational freightdata, models precise CO2 emission profiles, and applies predictive machine learning algorithms to replace guesswork with **Data-Driven Prescriptive Optimization**.":
    "Eine Softwareplattform wird entwickelt, die einen zentralen, datengetriebenen **Logistik-Kontrollturm** bereitstellt, der Rohdaten aus dem Frachtbetrieb aufnimmt, praezise CO2-Emissionsprofile modelliert und prädiktive Machine-Learning-Algorithmen einsetzt, um Schaedenschatzung durch **datengetriebene praskriptive Optimierung** zu ersetzen.",
    # Section 1.3
    "The platform provides a full-stack dashboard that allows logistics managers to:":
    "Die Plattform bietet ein Full-Stack-Dashboard, das Logistikmanagern folgende Moeglichkeiten bietet:",
    "- **Monitor real-time KPIs**: total CO2 emissions, distance traveled, orders fulfilled, and estimated carbon savings":
    "- **Echtzeit-KPIs ueberwachen**: Gesamte CO2-Emissionen, zurueckgelegte Strecke, erfuellte Auftraege und geschaetzte CO2-Einsparungen",
    "- **Manage fleet vehicles**: view all vehicles by type (electric, combustion, hybrid), track utilization, and filter by status":
    "- **Flottenfahrzeuge verwalten**: Alle Fahrzeuge nach Typ (Elektro, Verbrenner, Hybrid) anzeigen, Auslastung verfolgen und nach Status filtern",
    "- **Plan and optimize orders**: browse freight orders, view geographic routes on an interactive map, and compare planning scenarios":
    "- **Auftraege planen und optimieren**: Frachtauftraege durchsuchen, geographische Routen auf interaktiver Karte anzeigen und Planungsszenarien vergleichen",
    "- **Receive ML-generated alerts**: get prioritized recommendations for load consolidation, vehicle replacement, and route merging":
    "- **ML-generierte Warnungen erhalten**: Priorisierte Empfehlungen fuer Ladungskonsolidierung, Fahrzeugersatz und Routenzusammenfuehrung erhalten",
    "- **Run what-if simulations**: test alternative vehicle assignments using Machine Learning to project CO2 savings and utilization improvements":
    "- **Szenario-Simulationen durchfuehren**: Alternative Fahrzeugzuweisungen mit ML testen, um CO2-Einsparungen und Verbesserungen der Auslastung vorherzusagen",
    "- **Visualize routes geospatially**: display all delivery routes on an interactive Leaflet map with utilization-coded polylines and stop markers":
    "- **Routen geospatial visualisieren**: Alle Liefer Routen auf einer interaktiven Leaflet-Karte mit auslastungscodierten Polylinien und Haltemarkierungen anzeigen",
    "The system includes a complete data pipeline: CSV data ingestion, PostgreSQL relational modeling, analytics fact table construction, Random Forest ML model training, and a FastAPI REST API serving predictions to a Next.js React frontend.":
    "Das System umfasst eine vollstaendige Datenpipelinie: CSV-Datenaufnahme, relationale PostgreSQL-Modellierung, Erstellung von Analyse-Fakttabellen, Random-Forest-ML-Modelltraining und eine FastAPI-REST-API, die Vorhersagen an ein Next.js-React-Frontend ausliefert.",
    # Section 2.1
    "Every person using the system receives one of the following roles:":
    "Jede Person, die das System verwendet, erhaelt eine der folgenden Rollen:",
    "- **ADMIN** -- Full access to all dashboard features, fleet management, order planning, simulation, and alerts":
    "- **ADMIN** -- Vollzugriff auf alle Dashboard-Funktionen, Flottenverwaltung, Auftragsplanung, Simulation und Warnungen",
    "- **DISPATCHER** -- Can view dashboards, manage orders, run simulations, and acknowledge alerts":
    "- **DISPATCHER** -- Kann Dashboards anzeigen, Auftraege verwalten, Simulationen durchfuehren und Warnungen quittieren",
    "- **VIEWER** -- Read-only access to dashboards and reports":
    "- **VIEWER** -- Nur-Lese-Zugriff auf Dashboards und Berichte",
    "Owners of the \"ADMIN\" role may manage fleet vehicles, configure system settings, and access all analytics. Owners of the \"DISPATCHER\" role may manage orders and run simulations. Owners of the \"VIEWER\" role may only view dashboards and reports.":
    "Inhaber der Rolle \"ADMIN\" koennen Flottenfahrzeuge verwalten, Systemeinstellungen konfigurieren und auf alle Analysen zugreifen. Inhaber der Rolle \"DISPATCHER\" koennen Auftraege verwalten und Simulationen durchfuehren. Inhaber der Rolle \"VIEWER\" koennen nur Dashboards und Berichte anzeigen.",
    # Dashboard
    "As a logistics manager, I want to view real-time Key Performance Indicators (CO2 emissions, distance traveled, orders fulfilled, cost savings) so that I can monitor fleet performance at a glance.":
    "Als Logistikmanager moechte ich Echtzeit-Kennzahlen (CO2-Emissionen, zurueckgelegte Strecke, erfuellte Auftraege, Kosteneinsparungen) einsehen, um die Flottenleistung auf einen Blick zu ueberwachen.",
    "As a logistics manager, I want to see a fleet overview showing vehicle counts by type, utilization percentages, and electric vs. combustion breakdown so that I can assess fleet health.":
    "Als Logistikmanager moechte ich eine Flottenuebersicht mit Fahrzeuganzahlen nach Typ, Auslastungsprozenten und Elektro-vs.-Verbrenner-Aufschluesselung sehen, um den Flottenzustand zu beurteilen.",
    "As a logistics manager, I want a quick stats panel showing total vehicles, active routes, pending orders, and an eco score so that I have immediate access to critical metrics.":
    "Als Logistikmanager moechte ich ein Schnellstatistik-Panel mit Gesamtfahrzeugen, aktiven Routen, ausstehenden Auftraegen und einem Oko-Score haben, um sofortigen Zugriff auf kritische Kennzahlen zu haben.",
    "As a logistics manager, I want the dashboard to auto-refresh every 30 seconds so that I always see current data without manual intervention.":
    "Als Logistikmanager moechte ich, dass das Dashboard alle 30 Sekunden automatisch aktualisiert wird, damit ich immer aktuelle Daten ohne manuelle Eingriffe sehe.",
    # Fleet
    "As a fleet manager, I want to view all vehicles in a searchable grid with ID, type, tonnage, fuel type, emissions, battery level (for EVs), location, and status so that I can manage the fleet efficiently.":
    "Als Flottenmanager moechte ich alle Fahrzeuge in einem durchsuchbaren Raster mit ID, Typ, Tonnage, Kraftstoffart, Emissionen, Batteriestand (fuer EVs), Standort und Status sehen, um die Flotte effizient zu verwalten.",
    "As a fleet manager, I want to filter vehicles by type, fuel type, and status so that I can quickly find specific vehicle categories.":
    "Als Flottenmanager moechte ich Fahrzeuge nach Typ, Kraftstoffart und Status filtern, um bestimmte Fahrzeugkategorien schnell zu finden.",
    "As a fleet manager, I want to see per-type statistics including vehicle count, average load ratio, and emission intensity (kg CO2/km) so that I can compare vehicle type performance.":
    "Als Flottenmanager moechte ich typspezifische Statistiken mit Fahrzeuganzahl, durchschnittlicher Auslastung und Emissionsintensitaet (kg CO2/km) sehen, um die Leistung verschiedener Fahrzeugtypen zu vergleichen.",
    "As a fleet manager, I want to see the count of electric vs combustion vehicles so that I can track the fleet's electrification progress.":
    "Als Flottenmanager moechte ich die Anzahl der Elektro- vs. Verbrennungsfahrzeuge sehen, um den Elektrifizierungsfortschritt der Flotte zu verfolgen.",
    # Orders
    "As a dispatcher, I want to view all freight orders with distance, CO2 emissions, load ratio, and assigned vehicle so that I can manage deliveries.":
    "Als Disponent moechte ich alle Frachtauftraege mit Distanz, CO2-Emissionen, Auslastung und zugewiesenem Fahrzeug einsehen, um Lieferungen zu verwalten.",
    "As a dispatcher, I want to click on an order to see a detailed modal with KPIs (emissions, distance, load ratio, assigned vehicle) and geographic stops so that I can inspect individual orders.":
    "Als Disponent moechte ich auf einen Auftrag klicken, um ein Detailmodalfenster mit KPIs (Emissionen, Distanz, Auslastung, zugewiesenes Fahrzeug) und geographischen Halten zu oeffnen, um einzelne Auftraege zu pruefen.",
    "As a dispatcher, I want to filter orders by status (All, Pending, Assigned, Completed) and by route health (Optimal, Moderate, Low) so that I can focus on specific order categories.":
    "Als Disponent moechte ich Auftraege nach Status (Alle, Ausstehend, Zugewiesen, Abgeschlossen) und Routenzustand (Optimal, Moderat, Niedrig) filtern, um mich auf bestimmte Auftragskategorien zu konzentrieren.",
    "As a dispatcher, I want to compare planning scenarios (Normal vs Eco Planning) showing duration, emissions, cost, and efficiency so that I can choose the optimal strategy.":
    "Als Disponent moechte ich Planungsszenarien (Normal vs. Oko-Planung) mit Dauer, Emissionen, Kosten und Effizienz vergleichen, um die optimale Strategie zu waehlen.",
    # Alerts
    "As a logistics manager, I want to receive ML-generated alerts for load consolidation, vehicle replacement, and route merging so that I can optimize operations proactively.":
    "Als Logistikmanager moechte ich ML-generierte Warnungen fuer Ladungskonsolidierung, Fahrzeugersatz und Routenzusammenfuehrung erhalten, um Betriebsablaeufe proaktiv zu optimieren.",
    "As a logistics manager, I want each alert to have a priority score (High, Medium, Low) so that I can address the most impactful issues first.":
    "Als Logistikmanager moechte ich, dass jede Warnung einen Prioritaetsscore (Hoch, Mittel, Niedrig) hat, um die wirkungs vollen Probleme zuerst anzugehen.",
    "As a logistics manager, I want a summary view showing counts of high-priority, sustainability, and optimization alerts so that I can gauge overall fleet optimization status.":
    "Als Logistikmanager moechte ich eine Zusammenfassungsansicht mit Anzahl der hochprioritaetigen, Nachhaltigkeits- und Optimierungswarnungen, um den Gesamtstatus der Flottenoptimierung zu erfassen.",
    "As a logistics manager, I want to be able to take action on alerts (e.g., navigate to order, dismiss) so that I can act on recommendations directly.":
    "Als Logistikmanager moechte ich Massnahmen zu Warnungen ergreifen koennen (z.B. zum Auftrag navigieren, verwerfen), um direkt auf Empfehlungen reagieren zu koennen.",
    # Simulator
    "As a dispatcher, I want to select an order and test alternative vehicle types to project CO2 savings and utilization changes using Machine Learning so that I can make data-driven vehicle assignment decisions.":
    "Als Disponent moechte ich einen Auftrag auswaehlen und alternative Fahrzeugtypen testen, um CO2-Einsparungen und Auslastungsaenderungen mit ML vorherzusagen, um datengetriebene Entscheidungen ueber Fahrzeugzuweisungen zu treffen.",
    "As a dispatcher, I want the simulator to show all available vehicle types with their descriptions, capacity, and fuel type so that I can choose the best alternative.":
    "Als Disponent moechte ich, dass der Simulator alle verfuegbaren Fahrzeugtypen mit Beschreibungen, Kapazitaet und Kraftstoffart anzeigt, um die beste Alternative auszuwaehlen.",
    "As a dispatcher, I want to see a side-by-side comparison of current vs simulated emissions and utilization so that I can visually assess the impact of changing vehicles.":
    "Als Disponent moechte ich einen Seiten-an-Seiten-Vergleich aktueller vs. simulierter Emissionen und Auslastung sehen, um die Auswirkungen einer Fahrzeugaenderung visuell zu beurteilen.",
    "As a dispatcher, I want a visual bar chart comparing current vs simulated CO2 emissions so that I can quickly understand the environmental impact.":
    "Als Disponent moechte ich ein visuelles Balkendiagramm zum Vergleich aktueller vs. simulierter CO2-Emissionen sehen, um die Umweltauswirkungen schnell zu verstehen.",
    "As a dispatcher, I want to open an order from the details modal directly in the full simulator panel so that I can seamlessly transition from inspection to simulation.":
    "Als Disponent moechte ich einen Auftrag aus dem Detailmodalfenster direkt im vollstaendigen Simulator-Panel oeffnen, um nahtlos vom Pruefen zur Simulation zu wechseln.",
    # Map
    "As a dispatcher, I want to view all delivery routes on an interactive Leaflet map with polylines and stop markers so that I can visualize the geographic distribution of orders.":
    "Als Disponent moechte ich alle Liefer routen auf einer interaktiven Leaflet-Karte mit Polylinien und Haltemarkierungen sehen, um die geographische Verteilung der Auftraege zu visualisieren.",
    "As a dispatcher, I want routes to be color-coded by utilization (green for optimal, amber for moderate, red for low) with dashed lines for empty miles so that I can identify inefficient routes at a glance.":
    "Als Disponent moechte ich, dass Routen nach Auslastung farbcodiert sind (gruen fuer optimal, gelb fuer moderat, rot fuer niedrig) mit gestrichelten Linien fuer Leerfahrten, um ineffiziente Routen auf einen Blick zu erkennen.",
    "As a dispatcher, I want to click on a route to highlight it and filter the orders table so that I can focus on specific routes.":
    "Als Disponent moechte ich auf eine Route klicken, um sie hervorzuheben und die Auftragstabelle zu filtern, um mich auf bestimmte Routen zu konzentrieren.",
    "As a dispatcher, I want numbered stop markers (green for origin, red for destination) with popup details (city, coordinates, utilization) so that I can inspect individual stops.":
    "Als Disponent moechte ich nummerierte Haltemarkierungen (gruen fuer Ursprung, rot fuer Ziel) mit Popup-Details (Stadt, Koordinaten, Auslastung) haben, um einzelne Halte zu pruefen.",
    "As a dispatcher, I want a clickable sidebar showing route health counts (Optimal, Moderate, Low) that filters both the map and orders table so that I can quickly focus on problematic routes.":
    "Als Disponent moechte ich eine klickbare Seitenleiste mit Routenzustaenden (Optimal, Moderat, Niedrig), die sowohl Karte als auch Auftragstabelle filtert, um mich schnell auf problematische Routen zu konzentrieren.",
    # Profile
    "As a user, I want to view my profile information (name, email, phone, role, department, access level) so that I can verify my account details.":
    "Als Benutzer moechte ich meine Profilinformationen (Name, E-Mail, Telefon, Rolle, Abteilung, Zugriffsebene) einsehen, um meine Kontodaten zu ueberpruefen.",
    "As a user, I want to configure notification preferences (email alerts, push notifications, weekly reports) so that I control how I receive updates.":
    "Als Benutzer moechte ich Benachrichtigungseinstellungen (E-Mail-Warnungen, Push-Benachrichtigungen, Wochenberichte) konfigurieren, um zu kontrollieren, wie ich Updates erhalte.",
    "As a user, I want to generate Weekly, Monthly, and Annual reports so that I can export fleet performance data.":
    "Als Benutzer moechte ich Woechentliche, Monatliche und Jaehrliche Berichte erstellen, um Flottdaten exportieren zu koennen.",
}

def set_cell_shading(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_border(cell, top=None, bottom=None, left=None, right=None):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}></w:tcBorders>')
    for side, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        if val:
            border = parse_xml(f'<w:{side} {nsdecls("w")} w:val="single" w:sz="{val}" w:space="0" w:color="000000"/>')
            tcBorders.append(border)
    tcPr.append(tcBorders)

def add_formatted_paragraph(doc, text, style_name='Normal', bold=False, italic=False, font_size=None, alignment=None, space_before=None, space_after=None, font_color=None):
    p = doc.add_paragraph()
    if style_name and style_name != 'Normal':
        p.style = doc.styles[style_name]
    run = p.add_run(text)
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    if font_size:
        run.font.size = Pt(font_size)
    if font_color:
        run.font.color.rgb = font_color
    if alignment is not None:
        p.alignment = alignment
    if space_before is not None:
        p.paragraph_format.space_before = Pt(space_before)
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p

def process_inline_formatting(paragraph, text):
    """Handle bold (**text**), italic (*text*), and code (`text`) inline."""
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*'):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
        else:
            paragraph.add_run(part)

def build_docx(md_path, output_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    content = translate_full(content)
    content = content.replace('\\newpage', '')

    # Apply content translations
    for en, de in CONTENT_DE.items():
        content = content.replace(en, de)

    doc = Document()

    # Default font
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(10)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(4)

    # Heading styles
    for level in range(1, 4):
        hs = doc.styles[f'Heading {level}']
        hs.font.name = 'Arial'
        hs.font.color.rgb = RGBColor(0, 0, 0)

    # Margins
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # ---- COVER PAGE ----
    for _ in range(6):
        doc.add_paragraph()

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run("GreenTrack Logistik-Kontrollturm")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0, 0, 0)

    # Horizontal line
    hr_p = doc.add_paragraph()
    hr_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = hr_p._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="000000"/></w:pBdr>')
    pPr.append(pBdr)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_before = Pt(12)
    run = sub_p.add_run("Unternehmensplattform fuer datengetriebene\nSupply-Chain- und CO2-Optimierung")
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(60, 60, 60)

    for _ in range(4):
        doc.add_paragraph()

    author_p = doc.add_paragraph()
    author_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = author_p.add_run("Erstellt von:")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(80, 80, 80)

    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_p.paragraph_format.space_before = Pt(6)
    run = name_p.add_run("Balaji Loganathan")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)

    doc.add_paragraph()
    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = date_p.add_run("August 2026")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(80, 80, 80)

    doc.add_page_break()

    # ---- PARSE MARKDOWN ----
    lines = content.split('\n')
    i = 0
    in_code = False
    code_lines = []
    in_table = False
    table_rows = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Code blocks
        if stripped.startswith('```') and not in_code:
            in_code = True
            code_lines = []
            i += 1
            continue
        elif stripped.startswith('```') and in_code:
            in_code = False
            code_text = '\n'.join(code_lines)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(code_text)
            run.font.name = 'Courier New'
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(30, 30, 30)
            i += 1
            continue
        elif in_code:
            code_lines.append(line)
            i += 1
            continue

        # Tables
        if stripped.startswith('|') and '|' in stripped[1:]:
            cells = [c.strip() for c in stripped.split('|')[1:-1]]
            if not (len(cells) == 1 and all(c in '-: ' for c in cells[0])):
                table_rows.append(cells)
            i += 1
            continue
        elif table_rows:
            if table_rows:
                num_cols = max(len(r) for r in table_rows)
                table = doc.add_table(rows=len(table_rows), cols=num_cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.style = 'Table Grid'
                for ri, row_data in enumerate(table_rows):
                    for ci, cell_text in enumerate(row_data):
                        if ci < num_cols:
                            cell = table.cell(ri, ci)
                            cell.text = ''
                            p = cell.paragraphs[0]
                            run = p.add_run(cell_text)
                            run.font.name = 'Arial'
                            run.font.size = Pt(9)
                            run.font.color.rgb = RGBColor(0, 0, 0)
                            if ri == 0:
                                run.bold = True
                                set_cell_shading(cell, "E0E0E0")
                            elif ri % 2 == 0:
                                set_cell_shading(cell, "F5F5F5")
                doc.add_paragraph()
            table_rows = []

        # Headings
        if stripped.startswith('# ') and not stripped.startswith('## '):
            heading_text = stripped[2:].strip()
            if heading_text == 'Inhaltsverzeichnis':
                h = doc.add_heading(heading_text, level=1)
                for run in h.runs:
                    run.font.color.rgb = RGBColor(0, 0, 0)
            else:
                doc.add_page_break()
                h = doc.add_heading(heading_text, level=1)
                for run in h.runs:
                    run.font.color.rgb = RGBColor(0, 0, 0)
        elif stripped.startswith('## '):
            h = doc.add_heading(stripped[3:].strip(), level=2)
            for run in h.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
        elif stripped.startswith('### '):
            h = doc.add_heading(stripped[4:].strip(), level=3)
            for run in h.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
        elif stripped.startswith('- ') or stripped.startswith('* '):
            bullet_text = stripped[2:]
            p = doc.add_paragraph(style='List Bullet')
            p.text = ''
            process_inline_formatting(p, bullet_text)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
        elif re.match(r'^\d+\.', stripped):
            p = doc.add_paragraph(style='List Number')
            p.text = ''
            process_inline_formatting(p, stripped)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
        elif stripped == '---':
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="999999"/></w:pBdr>')
            pPr.append(pBdr)
        elif stripped == '':
            doc.add_paragraph()
        else:
            p = doc.add_paragraph()
            process_inline_formatting(p, stripped)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)

        i += 1

    # Flush remaining table
    if table_rows:
        num_cols = max(len(r) for r in table_rows)
        table = doc.add_table(rows=len(table_rows), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Table Grid'
        for ri, row_data in enumerate(table_rows):
            for ci, cell_text in enumerate(row_data):
                if ci < num_cols:
                    cell = table.cell(ri, ci)
                    cell.text = ''
                    p = cell.paragraphs[0]
                    run = p.add_run(cell_text)
                    run.font.name = 'Arial'
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(0, 0, 0)
                    if ri == 0:
                        run.bold = True
                        set_cell_shading(cell, "E0E0E0")
                    elif ri % 2 == 0:
                        set_cell_shading(cell, "F5F5F5")

    doc.save(output_path)
    return output_path

if __name__ == '__main__':
    md_path = r"D:\projects\Makeathon\docs\PROJECT_DOCUMENTATION.md"
    out = build_docx(md_path, OUTPUT_PATH)
    print(f"DOCX generated: {out}")
